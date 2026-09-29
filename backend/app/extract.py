"""Plain-text extraction for files placed in a 資料室."""

from pathlib import Path

MAX_TEXT_CHARS = 2_000_000

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".log", ".html", ".htm", ".xml", ".yaml", ".yml"}
EDITABLE_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".tsv"}
EXTRACTABLE_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx", ".xlsx", ".pptx"}


def decode_text(data: bytes) -> str:
    """Decode bytes from files made on Japanese Windows as well as UTF-8."""
    for encoding in ("utf-8-sig", "cp932"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(path: Path) -> str:
    """Return the file's text, or "" for formats we do not read. Raises on unreadable files."""
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        text = decode_text(path.read_bytes())
    elif suffix == ".pdf":
        text = _pdf(path)
    elif suffix == ".docx":
        text = _docx(path)
    elif suffix == ".xlsx":
        text = _xlsx(path)
    elif suffix == ".pptx":
        text = _pptx(path)
    else:
        return ""
    # PostgreSQL text columns cannot hold NUL characters.
    return text.replace("\x00", "")[:MAX_TEXT_CHARS]


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if v is None else str(v) for v in row]
                if any(values):
                    parts.append("\t".join(values))
    finally:
        workbook.close()
    return "\n".join(parts)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts.append(f"# スライド {number}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)
